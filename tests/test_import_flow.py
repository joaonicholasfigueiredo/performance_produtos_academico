from __future__ import annotations

import unittest
from io import BytesIO
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

from app.config import EXPECTED_COLUMNS
from services import analytics, audit, exporter, importer
from repo import database


def build_workbook(second_production: float = 10) -> bytes:
    frame = pd.DataFrame(
        [
            {
                "CENTRAL": "Central A",
                "COOP": "0001",
                "PA": "01",
                "PRODUTO": "Produto A",
                "ID_PRODUTO": "P001",
                "META_MENSAL": "5.70",
                "PRODUCAO_MENSAL": None,
                "META_TRIMESTRAL": "1.234,56",
                "PRODUCAO_TRIMESTRAL": 1200,
                "META_SEMESTRAL": 2000,
                "PRODUCAO_SEMESTRAL": 1900,
                "META_ANUAL": 4000,
                "PRODUCAO_ANUAL": 3900,
                "META_NO_PERIODO": 2000,
                "PRODUCAO_NO_PERIODO": "2,50",
                "MÊS/ANO": "07/2026",
            },
            {
                "CENTRAL": "Central A",
                "COOP": "0001",
                "PA": "02",
                "PRODUTO": "Produto A",
                "ID_PRODUTO": "P001",
                "META_MENSAL": "4,30",
                "PRODUCAO_MENSAL": second_production,
                "META_TRIMESTRAL": 10,
                "PRODUCAO_TRIMESTRAL": 10,
                "META_SEMESTRAL": 10,
                "PRODUCAO_SEMESTRAL": 10,
                "META_ANUAL": 10,
                "PRODUCAO_ANUAL": 10,
                "META_NO_PERIODO": 10,
                "PRODUCAO_NO_PERIODO": 10,
                "MÊS/ANO": "07/2026",
            },
        ],
        columns=EXPECTED_COLUMNS,
    )
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False)
    return output.getvalue()


class MemoryStorage:
    def __init__(self) -> None:
        self.frame = pd.DataFrame()
        self.events: list[dict] = []
        self.pending: dict[str, pd.DataFrame] = {}

    def period_exists(self, period: pd.Timestamp) -> bool:
        if self.frame.empty:
            return False
        return bool((pd.to_datetime(self.frame[importer.PERIOD_COLUMN]) == period).any())

    def load_period_data(self, period: pd.Timestamp) -> pd.DataFrame:
        if self.frame.empty:
            return pd.DataFrame(columns=database.ALL_COLUMNS)
        mask = pd.to_datetime(self.frame[importer.PERIOD_COLUMN]) == period
        return self.frame.loc[mask].copy()

    def replace_period_data(
        self,
        incoming: pd.DataFrame,
        event: dict | None = None,
        pending_token: str | None = None,
    ) -> int:
        period = pd.Timestamp(incoming[importer.PERIOD_COLUMN].iloc[0])
        if self.frame.empty:
            self.frame = incoming.copy()
        else:
            mask = pd.to_datetime(self.frame[importer.PERIOD_COLUMN]) == period
            remaining = self.frame.loc[~mask]
            self.frame = (
                incoming.copy()
                if remaining.empty
                else pd.concat([remaining, incoming], ignore_index=True)
            )
        if event is not None:
            self.events.append(event)
        if pending_token is not None:
            self.pending.pop(pending_token, None)
        return len(self.frame)

    def load_product_data(self, columns=None) -> pd.DataFrame:
        selected = list(columns or database.ALL_COLUMNS)
        if self.frame.empty:
            return pd.DataFrame(columns=selected)
        return self.frame[selected].copy()

    def load_log_events(self) -> list[dict]:
        return list(reversed(self.events))

    def save_pending_import(self, token: str, frame: pd.DataFrame) -> None:
        self.pending[token] = frame.copy()

    def load_pending_import(self, token: str) -> pd.DataFrame:
        frame = self.pending.get(token)
        return (
            frame.copy()
            if frame is not None
            else pd.DataFrame(columns=EXPECTED_COLUMNS)
        )

    def delete_pending_import(self, token: str) -> bool:
        return self.pending.pop(token, None) is not None

    def count_pending_imports(self) -> int:
        return len(self.pending)

    def storage_status(self) -> dict:
        periods = (
            sorted(pd.to_datetime(self.frame[importer.PERIOD_COLUMN]).drop_duplicates())
            if not self.frame.empty
            else []
        )
        return {"exists": not self.frame.empty, "total_rows": len(self.frame), "periods": periods}

    def drop_storage(self) -> None:
        self.frame = pd.DataFrame()
        self.events = []
        self.pending = {}


class ImportFlowTest(unittest.TestCase):
    def setUp(self) -> None:
        self.storage = MemoryStorage()

        self.patches = [
            patch.object(database, "period_exists", self.storage.period_exists),
            patch.object(database, "load_period_data", self.storage.load_period_data),
            patch.object(database, "replace_period_data", self.storage.replace_period_data),
            patch.object(database, "load_product_data", self.storage.load_product_data),
            patch.object(database, "load_log_events", self.storage.load_log_events),
            patch.object(database, "save_pending_import", self.storage.save_pending_import),
            patch.object(database, "load_pending_import", self.storage.load_pending_import),
            patch.object(database, "delete_pending_import", self.storage.delete_pending_import),
            patch.object(database, "count_pending_imports", self.storage.count_pending_imports),
            patch.object(database, "storage_status", self.storage.storage_status),
            patch.object(database, "drop_storage", self.storage.drop_storage),
        ]
        for active_patch in self.patches:
            active_patch.start()

    def tearDown(self) -> None:
        for active_patch in reversed(self.patches):
            active_patch.stop()

    def test_analyze_confirm_deduplicate_and_export(self) -> None:
        workbook = build_workbook()
        analysis = importer.analyze_upload("julho.xlsx", workbook)
        self.assertIn(analysis.token, self.storage.pending)

        self.assertEqual(analysis.period_label, "Julho/2026")
        self.assertEqual(analysis.rows, 2)
        self.assertGreaterEqual(analysis.blank_values_filled, 1)
        self.assertEqual(analysis.preview[0]["META_MENSAL"], 5.7)
        self.assertEqual(analysis.preview[0]["PRODUCAO_MENSAL"], 0.0)
        self.assertEqual(analysis.preview[0]["META_TRIMESTRAL"], 1234.56)

        first = importer.confirm_import(analysis.token, analysis.file_name)
        self.assertNotIn(analysis.token, self.storage.pending)
        self.assertEqual(first["inserted_rows"], 2)
        self.assertEqual(first["total_rows"], 2)

        periods = analytics.available_periods()
        self.assertEqual(periods["default_start"], "2026-07")
        self.assertEqual(periods["default_unit"], "SINGULAR")
        self.assertEqual(
            [unit["value"] for unit in periods["units"]],
            ["SINGULAR", "01", "02"],
        )
        tracking = analytics.accumulated_production("2026-07", "2026-07")
        self.assertEqual(tracking["summary"]["products"], 1)
        self.assertEqual(tracking["summary"]["adherent"], 1)
        self.assertEqual(tracking["products"][0]["produto"], "Produto A")
        self.assertAlmostEqual(tracking["products"][0]["meta_acumulada"], 10)
        self.assertAlmostEqual(tracking["products"][0]["producao_acumulada"], 10)

        pa_tracking = analytics.accumulated_production("2026-07", "2026-07", "01")
        self.assertEqual(pa_tracking["unit"]["label"], "PA 01")
        self.assertEqual(pa_tracking["summary"]["not_adherent"], 1)
        self.assertAlmostEqual(pa_tracking["products"][0]["meta_acumulada"], 5.7)
        self.assertAlmostEqual(pa_tracking["products"][0]["producao_acumulada"], 0)

        updated_workbook = build_workbook(second_production=15)
        second_analysis = importer.analyze_upload("julho_corrigido.xlsx", updated_workbook)
        self.assertTrue(second_analysis.period_already_exists)
        second = importer.confirm_import(second_analysis.token, second_analysis.file_name)
        self.assertEqual(second["action"], "updated")
        self.assertEqual(second["inserted_rows"], 0)
        self.assertEqual(second["updated_rows"], 1)
        self.assertEqual(second["removed_rows"], 0)
        self.assertEqual(second["changed_rows"], 1)
        self.assertEqual(second["total_rows"], 2)
        self.assertIsNotNone(second["log_file"])

        logs = audit.query_update_logs()
        self.assertEqual(logs["summary"]["updates"], 1)
        self.assertEqual(logs["summary"]["changes"], 1)
        self.assertEqual(logs["entries"][0]["produto"], "Produto A")
        self.assertEqual(logs["entries"][0]["pa"], "02")
        self.assertEqual(logs["entries"][0]["type"], "ATUALIZADO")
        self.assertEqual(logs["entries"][0]["before"]["PRODUCAO_MENSAL"], 10)
        self.assertEqual(logs["entries"][0]["after"]["PRODUCAO_MENSAL"], 15)

        output_content, output_name = exporter.export_base()
        self.assertIsInstance(output_content, bytes)
        self.assertTrue(output_name.endswith(".xlsx"))

        exported = pd.read_excel(BytesIO(output_content))
        self.assertEqual(list(exported.columns), EXPECTED_COLUMNS)
        self.assertEqual(len(exported), 2)

        pending = importer.analyze_upload("julho_pendente.xlsx", workbook)
        self.assertIn(pending.token, self.storage.pending)
        wiped = importer.wipe_base()
        self.assertTrue(wiped["deleted"])
        self.assertEqual(wiped["removed_rows"], 2)
        self.assertEqual(wiped["removed_periods"], 1)
        self.assertEqual(wiped["removed_pending_imports"], 1)
        self.assertFalse(self.storage.pending)
        self.assertFalse(importer.base_status()["exists"])

    def test_tracking_report_has_four_management_sheets(self) -> None:
        rows = []
        for period, production in (
            (pd.Timestamp("2026-01-01"), 90),
            (pd.Timestamp("2026-02-01"), 120),
        ):
            for unit in ("0", "01"):
                rows.append(
                    {
                        "MÊS/ANO": period,
                        "PA": unit,
                        "PRODUTO": "COBRANCA",
                        "META_MENSAL": 100,
                        "PRODUCAO_MENSAL": production,
                        "META_ANUAL": 1200,
                        "PRODUCAO_ANUAL": production,
                    }
                )
        self.storage.frame = pd.DataFrame(rows)

        content, name = exporter.export_tracking_table(
            "SINTETICO",
            "2026-01",
            "2026-02",
            "01",
            "NO_PERIODO",
            "FOCO",
        )

        self.assertIsInstance(content, bytes)
        self.assertTrue(name.startswith("relatorio_acompanhamento_2026-01_a_2026-02_"))
        workbook = load_workbook(BytesIO(content), data_only=True)
        self.addCleanup(workbook.close)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Produção Mensal",
                "Resumo Consolidado",
                "Desempenho 19 Produtos",
                "Desempenho 7 Produtos",
            ],
        )

        monthly = workbook["Produção Mensal"]
        self.assertEqual(monthly["A1"].value, "ACOMPANHAMENTO DE PRODUTOS")
        self.assertEqual(monthly["A2"].value, "PRODUÇÃO MENSAL POR PRODUTO")
        self.assertIn("JANEIRO/2026 A FEVEREIRO/2026", monthly["A3"].value)
        self.assertEqual(monthly["A5"].value, "PRODUÇÃO EM VOLUME — JANEIRO/2026")
        self.assertEqual(
            [monthly.cell(6, column).value for column in range(1, 6)],
            ["PRODUTO", "META", "PRODUÇÃO", "%", "GAP"],
        )
        self.assertEqual(monthly["A10"].value, "PRODUÇÃO EM VOLUME — FEVEREIRO/2026")
        self.assertEqual(monthly["B7"].number_format, "R$ #,##0;[Red]-R$ #,##0")
        self.assertEqual(monthly["D7"].number_format, '0.0"%"')
        self.assertEqual(monthly["A1"].fill.fgColor.rgb, "00003641")
        self.assertEqual(monthly["A4"].fill.fgColor.rgb, "00C9D200")
        self.assertEqual(monthly["A5"].fill.fgColor.rgb, "00003641")
        self.assertEqual(monthly["A6"].fill.fgColor.rgb, "0000AE9D")
        self.assertEqual(monthly["D7"].fill.fgColor.rgb, "00FFF0ED")
        self.assertIsNone(monthly.freeze_panes)
        self.assertFalse(monthly.sheet_view.showGridLines)

        consolidated = workbook["Resumo Consolidado"]
        self.assertEqual(
            consolidated["A5"].value,
            "PRODUÇÃO ACUMULADA — JANEIRO/2026 A FEVEREIRO/2026",
        )
        self.assertEqual(
            consolidated["A10"].value,
            "PRODUÇÃO ANUAL — ÚLTIMA BASE DISPONÍVEL: FEVEREIRO/2026",
        )

        general = workbook["Desempenho 19 Produtos"]
        self.assertEqual(general["A5"].value, "DESEMPENHO 19 PRODUTOS — SINGULAR")
        self.assertEqual(
            [general.cell(6, column).value for column in range(1, 7)],
            [
                "PRODUTO",
                "JANEIRO/2026",
                "FEVEREIRO/2026",
                "METAS NO PERÍODO",
                "METAS ATINGIDAS",
                "% APROVEITAMENTO",
            ],
        )
        self.assertEqual(general["A10"].value, "DESEMPENHO 19 PRODUTOS — AGÊNCIAS")
        self.assertEqual(
            [general.cell(11, column).value for column in range(1, 7)],
            [
                "AGÊNCIA",
                "JANEIRO/2026",
                "FEVEREIRO/2026",
                "METAS NO PERÍODO",
                "METAS ATINGIDAS",
                "% APROVEITAMENTO",
            ],
        )
        self.assertEqual([general["B7"].value, general["C7"].value], [0, 1])
        self.assertEqual(general["D7"].value, 2)
        self.assertEqual(general["E7"].value, 1)
        self.assertEqual(general["F7"].value, 50)
        self.assertEqual(general["B7"].fill.fgColor.rgb, "00F0F4F2")
        self.assertEqual(general["C7"].fill.fgColor.rgb, "00E6F6EB")
        self.assertEqual(general["D12"].value, 36)
        self.assertEqual(general["E12"].value, 1)
        self.assertAlmostEqual(general["F12"].value, 1 / 36 * 100)
        self.assertEqual(general["F12"].number_format, '0.0"%"')
        self.assertEqual(len(general.conditional_formatting), 2)

        focus = workbook["Desempenho 7 Produtos"]
        self.assertEqual(focus["A5"].value, "DESEMPENHO 7 PRODUTOS — SINGULAR")
        self.assertEqual(focus["A10"].value, "DESEMPENHO 7 PRODUTOS — AGÊNCIAS")
        self.assertEqual(focus["D7"].value, 2)
        self.assertEqual(focus["E7"].value, 1)
        self.assertEqual(focus["F7"].value, 50)
        self.assertEqual(focus["D12"].value, 14)
        self.assertEqual(focus["E12"].value, 1)
        self.assertAlmostEqual(focus["F12"].value, 1 / 14 * 100)

    def test_venda_nova_uses_the_product_for_the_selected_scope(self) -> None:
        frame = pd.DataFrame(
            [
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": "00",
                    "PRODUTO": "VENDA_NOVA_SINGULAR",
                    "META_MENSAL": 100,
                    "PRODUCAO_MENSAL": 110,
                    "META_ANUAL": 100,
                    "PRODUCAO_ANUAL": 110,
                    "META_NO_PERIODO": 80,
                    "PRODUCAO_NO_PERIODO": 88,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": "01",
                    "PRODUTO": "VENDA_NOVA_PA",
                    "META_MENSAL": 40,
                    "PRODUCAO_MENSAL": 35,
                    "META_ANUAL": 40,
                    "PRODUCAO_ANUAL": 35,
                    "META_NO_PERIODO": 30,
                    "PRODUCAO_NO_PERIODO": 28,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": "02",
                    "PRODUTO": "VENDA_NOVA_PA",
                    "META_MENSAL": 50,
                    "PRODUCAO_MENSAL": 55,
                    "META_ANUAL": 50,
                    "PRODUCAO_ANUAL": 55,
                    "META_NO_PERIODO": 35,
                    "PRODUCAO_NO_PERIODO": 39,
                },
            ]
        )
        self.storage.frame = frame

        singular = analytics.accumulated_production(
            "2026-01", "2026-01", "SINGULAR"
        )
        self.assertEqual(
            [product["produto"] for product in singular["products"]],
            ["VENDA_NOVA_SINGULAR"],
        )
        self.assertEqual(singular["products"][0]["meta_acumulada"], 100)
        self.assertEqual(singular["products"][0]["producao_acumulada"], 110)

        pa = analytics.accumulated_production("2026-01", "2026-01", "01")
        self.assertEqual(
            [product["produto"] for product in pa["products"]],
            ["VENDA_NOVA_PA"],
        )
        self.assertAlmostEqual(pa["products"][0]["meta_acumulada"], 100 * 40 / 90)
        self.assertEqual(pa["products"][0]["producao_acumulada"], 35)

        analytic_singular = analytics.monthly_product_evolution(
            "2026-01", "2026-01", "SINGULAR"
        )
        analytic_pa = analytics.monthly_product_evolution(
            "2026-01", "2026-01", "01"
        )
        self.assertEqual(
            [product["produto"] for product in analytic_singular["products"]],
            ["VENDA_NOVA_SINGULAR"],
        )
        self.assertEqual(
            [product["produto"] for product in analytic_pa["products"]],
            ["VENDA_NOVA_PA"],
        )

        annual = analytics.accumulated_production(
            "2026-01", "2026-01", "SINGULAR", "ANUAL"
        )
        period = analytics.accumulated_production(
            "2026-01", "2026-01", "01", "NO_PERIODO"
        )
        self.assertEqual(annual["calculation"]["production_field"], "PRODUCAO_ANUAL")
        self.assertEqual(annual["products"][0]["meta_acumulada"], 100)
        self.assertEqual(annual["products"][0]["producao_acumulada"], 110)
        self.assertEqual(period["calculation"]["production_field"], "PRODUCAO_MENSAL")
        self.assertAlmostEqual(period["products"][0]["meta_acumulada"], 100 * 40 / 90)
        self.assertEqual(period["products"][0]["producao_acumulada"], 35)

        focus_singular = analytics.accumulated_production(
            "2026-01", "2026-01", "SINGULAR", "ANUAL", "FOCO"
        )
        focus_pa = analytics.monthly_product_evolution(
            "2026-01", "2026-01", "01", "FOCO"
        )
        self.assertEqual(focus_singular["scope"]["value"], "FOCO")
        self.assertEqual(
            [product["produto"] for product in focus_singular["products"]],
            ["VENDA_NOVA_PA"],
        )
        self.assertEqual(focus_singular["products"][0]["meta_acumulada"], 90)
        self.assertEqual(focus_singular["products"][0]["producao_acumulada"], 90)
        self.assertEqual(
            [product["produto"] for product in focus_pa["products"]],
            ["VENDA_NOVA_PA"],
        )

    def test_venda_nova_pa_target_is_not_reduced_or_adjusted_without_singular(self) -> None:
        frame = pd.DataFrame(
            [
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": "00",
                    "PRODUTO": "VENDA_NOVA_SINGULAR",
                    "META_MENSAL": 80,
                    "PRODUCAO_MENSAL": 0,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": "01",
                    "PRODUTO": "VENDA_NOVA_PA",
                    "META_MENSAL": 40,
                    "PRODUCAO_MENSAL": 0,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": "02",
                    "PRODUTO": "VENDA_NOVA_PA",
                    "META_MENSAL": 50,
                    "PRODUCAO_MENSAL": 0,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-02-01"),
                    "PA": "01",
                    "PRODUTO": "VENDA_NOVA_PA",
                    "META_MENSAL": 60,
                    "PRODUCAO_MENSAL": 0,
                },
            ]
        )

        adjusted = analytics._apply_business_allocations(frame)
        january_pa = adjusted[
            (adjusted["MÊS/ANO"] == pd.Timestamp("2026-01-01"))
            & (adjusted["PRODUTO"] == "VENDA_NOVA_PA")
        ]
        february_pa = adjusted[
            (adjusted["MÊS/ANO"] == pd.Timestamp("2026-02-01"))
            & (adjusted["PRODUTO"] == "VENDA_NOVA_PA")
        ]

        self.assertEqual(january_pa["META_MENSAL"].tolist(), [40, 50])
        self.assertEqual(february_pa["META_MENSAL"].tolist(), [60])

    def test_unit_summary_includes_headquarters_pa_zero(self) -> None:
        self.storage.frame = pd.DataFrame(
            [
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": "0",
                    "PRODUTO": "COBRANCA",
                    "META_MENSAL": 100,
                    "PRODUCAO_MENSAL": 110,
                    "META_ANUAL": 100,
                    "PRODUCAO_ANUAL": 110,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": "01",
                    "PRODUTO": "COBRANCA",
                    "META_MENSAL": 100,
                    "PRODUCAO_MENSAL": 90,
                    "META_ANUAL": 100,
                    "PRODUCAO_ANUAL": 90,
                },
            ]
        )

        summary = analytics.unit_performance_summary(
            "2026-01", "2026-01", "NO_PERIODO", "GERAL"
        )
        headquarters = next(unit for unit in summary["units"] if unit["pa"] == "0")
        self.assertEqual(headquarters["code"], "00")
        self.assertEqual(headquarters["name"], "SEDE")
        self.assertEqual(headquarters["label"], "PA 00 — SEDE")
        self.assertEqual(headquarters["products"], 18)
        self.assertEqual(headquarters["adherent"], 1)

    def test_unit_summary_counts_monthly_opportunities_and_reallocates_life_insurance(self) -> None:
        rows = []
        for period, cobranca_production, insurance_production in (
            ("2026-01-01", 110, 120),
            ("2026-02-01", 90, 100),
        ):
            rows.extend(
                [
                    {
                        "MÊS/ANO": pd.Timestamp(period),
                        "PA": "0",
                        "PRODUTO": "COBRANCA",
                        "META_MENSAL": 100,
                        "PRODUCAO_MENSAL": cobranca_production,
                        "META_ANUAL": 200,
                        "PRODUCAO_ANUAL": cobranca_production,
                    },
                    {
                        "MÊS/ANO": pd.Timestamp(period),
                        "PA": "0",
                        "PRODUTO": "SEGURO_VIDA",
                        "META_MENSAL": 100,
                        "PRODUCAO_MENSAL": insurance_production,
                        "META_ANUAL": 200,
                        "PRODUCAO_ANUAL": insurance_production,
                    },
                    {
                        "MÊS/ANO": pd.Timestamp(period),
                        "PA": "01",
                        "PRODUTO": "COBRANCA",
                        "META_MENSAL": 100,
                        "PRODUCAO_MENSAL": 100,
                        "META_ANUAL": 200,
                        "PRODUCAO_ANUAL": 100,
                    },
                ]
            )
        self.storage.frame = pd.DataFrame(rows)

        summary = analytics.unit_performance_summary(
            "2026-01", "2026-02", "NO_PERIODO", "GERAL"
        )
        units = {unit["code"]: unit for unit in summary["units"]}

        self.assertEqual(summary["scope"]["products_per_month"], 18)
        self.assertEqual(summary["scope"]["months"], 2)
        self.assertEqual(summary["scope"]["target"], 36)
        self.assertEqual(units["00"]["products"], 36)
        self.assertEqual(units["00"]["adherent"], 1)
        self.assertEqual(units["01"]["adherent"], 2)
        self.assertNotIn("89", units)
        self.assertEqual(
            units["00"]["monthly"][0]["achieved_products"],
            ["COBRANCA"],
        )
        self.assertEqual(units["00"]["monthly"][1]["achieved_products"], [])

        headquarters = analytics.accumulated_production(
            "2026-01", "2026-02", "0", "NO_PERIODO"
        )
        uad = analytics.accumulated_production(
            "2026-01", "2026-02", "89", "NO_PERIODO"
        )
        self.assertNotIn(
            "SEGURO_VIDA_ARRECADACAO",
            [product["produto"] for product in headquarters["products"]],
        )
        self.assertIn(
            "SEGURO_VIDA_ARRECADACAO",
            [product["produto"] for product in uad["products"]],
        )

        annual_summary = analytics.unit_performance_summary(
            "2026-01", "2026-02", "ANUAL", "GERAL"
        )
        annual_units = {unit["code"]: unit for unit in annual_summary["units"]}
        self.assertEqual(annual_summary["scope"]["months"], 12)
        self.assertEqual(annual_summary["scope"]["period_months"], 2)
        self.assertEqual(annual_summary["scope"]["target"], 216)
        self.assertEqual(annual_units["00"]["products"], 216)
        self.assertEqual(annual_units["00"]["adherent"], 1)

        annual_focus = analytics.unit_performance_summary(
            "2026-01", "2026-02", "ANUAL", "FOCO"
        )
        focus_units = {unit["code"]: unit for unit in annual_focus["units"]}
        self.assertEqual(annual_focus["scope"]["target"], 84)
        self.assertEqual(focus_units["00"]["products"], 84)
        self.assertEqual(focus_units["00"]["adherent"], 1)

    def test_agency_cards_group_products_and_keep_credit_latest_balance(self) -> None:
        self.storage.frame = pd.DataFrame(
            [
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": "01",
                    "PRODUTO": "COBRANCA",
                    "META_MENSAL": 100,
                    "PRODUCAO_MENSAL": 90,
                    "META_ANUAL": 100,
                    "PRODUCAO_ANUAL": 90,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-02-01"),
                    "PA": "01",
                    "PRODUTO": "COBRANCA",
                    "META_MENSAL": 100,
                    "PRODUCAO_MENSAL": 120,
                    "META_ANUAL": 200,
                    "PRODUCAO_ANUAL": 210,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": "01",
                    "PRODUTO": "CREDITO_PESSOAL",
                    "META_MENSAL": 50,
                    "PRODUCAO_MENSAL": 40,
                    "META_ANUAL": 490,
                    "PRODUCAO_ANUAL": 460,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-02-01"),
                    "PA": "01",
                    "PRODUTO": "CREDITO_PESSOAL",
                    "META_MENSAL": 50,
                    "PRODUCAO_MENSAL": 45,
                    "META_ANUAL": 500,
                    "PRODUCAO_ANUAL": 450,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-02-01"),
                    "PA": "04",
                    "PRODUTO": "COBRANCA",
                    "META_MENSAL": 80,
                    "PRODUCAO_MENSAL": 90,
                    "META_ANUAL": 160,
                    "PRODUCAO_ANUAL": 170,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-02-01"),
                    "PA": "89",
                    "PRODUTO": "COBRANCA",
                    "META_MENSAL": 80,
                    "PRODUCAO_MENSAL": 90,
                    "META_ANUAL": 160,
                    "PRODUCAO_ANUAL": 170,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-02-01"),
                    "PA": "999",
                    "PRODUTO": "COBRANCA",
                    "META_MENSAL": 80,
                    "PRODUCAO_MENSAL": 90,
                    "META_ANUAL": 160,
                    "PRODUCAO_ANUAL": 170,
                },
            ]
        )

        result = analytics.agency_accumulated_production(
            "2026-01", "2026-02", "ANUAL"
        )

        self.assertEqual(result["summary"]["units"], 2)
        self.assertEqual([unit["code"] for unit in result["units"]], ["01", "04"])
        periods = analytics.available_periods()
        self.assertEqual(
            [unit["value"] for unit in periods["units"]],
            ["SINGULAR", "01", "04"],
        )
        singular = analytics.accumulated_production(
            "2026-01", "2026-02", "SINGULAR", "NO_PERIODO"
        )
        singular_products = {
            product["produto"]: product for product in singular["products"]
        }
        self.assertEqual(singular_products["COBRANCA"]["producao_acumulada"], 480)
        first_unit = result["units"][0]
        self.assertEqual(first_unit["name"], "PA 01")
        products = {product["produto"]: product for product in first_unit["products"]}
        self.assertEqual(products["COBRANCA"]["meta_acumulada"], 200)
        self.assertEqual(products["COBRANCA"]["producao_acumulada"], 210)
        self.assertEqual(products["CREDITO_PESSOAL"]["meta_acumulada"], 500)
        self.assertEqual(products["CREDITO_PESSOAL"]["producao_acumulada"], 450)
        self.assertEqual(
            products["CREDITO_PESSOAL"]["regra_calculo"],
            "ultimo_saldo_anual",
        )
        period_result = analytics.agency_accumulated_production(
            "2026-01", "2026-02", "NO_PERIODO"
        )
        period_products = {
            product["produto"]: product
            for product in period_result["units"][0]["products"]
        }
        self.assertEqual(period_result["calculation"]["value"], "NO_PERIODO")
        self.assertEqual(period_products["CREDITO_PESSOAL"]["meta_acumulada"], 50)
        self.assertEqual(period_products["CREDITO_PESSOAL"]["producao_acumulada"], 45)
        self.assertEqual(
            period_products["CREDITO_PESSOAL"]["regra_calculo"],
            "saldo_mensal_nao_acumulado",
        )

    def test_agency_score_keeps_eighteen_opportunities_when_a_product_has_no_target(self) -> None:
        self.storage.frame = pd.DataFrame(
            [
                {
                    "MÊS/ANO": pd.Timestamp("2026-08-01"),
                    "PA": "01",
                    "PRODUTO": "COBRANCA",
                    "META_MENSAL": 100,
                    "PRODUCAO_MENSAL": 120,
                    "META_ANUAL": 1200,
                    "PRODUCAO_ANUAL": 900,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-08-01"),
                    "PA": "01",
                    "PRODUTO": "PREVIDENCIA_VGBL",
                    "META_MENSAL": 0,
                    "PRODUCAO_MENSAL": 1,
                    "META_ANUAL": 5,
                    "PRODUCAO_ANUAL": 4,
                },
            ]
        )

        result = analytics.agency_accumulated_production("2026-08", "2026-08")
        summary = result["units"][0]["summary"]

        self.assertEqual(summary["adherent"], 1)
        self.assertEqual(summary["with_target"], 18)
        self.assertEqual(summary["not_adherent"], 17)
        self.assertEqual(summary["without_target"], 17)
        self.assertAlmostEqual(summary["adherence_rate"], 100 / 18)

    def test_credit_products_use_latest_annual_balance(self) -> None:
        frame = pd.DataFrame(
            [
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": "01",
                    "PRODUTO": "FUNDOS",
                    "META_MENSAL": 10,
                    "PRODUCAO_MENSAL": 8,
                    "META_ANUAL": 120,
                    "PRODUCAO_ANUAL": 8,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-02-01"),
                    "PA": "01",
                    "PRODUTO": "FUNDOS",
                    "META_MENSAL": 10,
                    "PRODUCAO_MENSAL": 12,
                    "META_ANUAL": 120,
                    "PRODUCAO_ANUAL": 20,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": "01",
                    "PRODUTO": "CREDITO_PESSOAL",
                    "META_MENSAL": 50,
                    "PRODUCAO_MENSAL": 45,
                    "META_ANUAL": 100,
                    "PRODUCAO_ANUAL": 90,
                },
                {
                    "MÊS/ANO": pd.Timestamp("2026-02-01"),
                    "PA": "01",
                    "PRODUTO": "CREDITO_PESSOAL",
                    "META_MENSAL": 60,
                    "PRODUCAO_MENSAL": 55,
                    "META_ANUAL": 110,
                    "PRODUCAO_ANUAL": 105,
                },
            ]
        )
        self.storage.frame = frame

        result = analytics.accumulated_production("2026-01", "2026-02", "SINGULAR")
        products = {item["produto"]: item for item in result["products"]}

        self.assertEqual(result["range"]["title"], "Produção (Janeiro a Fevereiro)")
        self.assertEqual(products["FUNDOS"]["meta_acumulada"], 20)
        self.assertEqual(products["FUNDOS"]["producao_acumulada"], 20)
        self.assertEqual(products["FUNDOS"]["regra_calculo"], "soma_mensal")
        self.assertEqual(products["CREDITO_PESSOAL"]["meta_acumulada"], 110)
        self.assertEqual(products["CREDITO_PESSOAL"]["producao_acumulada"], 105)

        self.assertEqual(products["CREDITO_PESSOAL"]["gap"], -5)
        self.assertEqual(
            products["CREDITO_PESSOAL"]["regra_calculo"],
            "ultimo_saldo_anual",
        )
        self.assertEqual(products["CREDITO_PESSOAL"]["periodo_referencia"], "2026-02")

        annual = analytics.accumulated_production(
            "2026-01", "2026-02", "SINGULAR", "ANUAL"
        )
        annual_products = {item["produto"]: item for item in annual["products"]}
        self.assertEqual(annual["calculation"]["reference_period"], "2026-02")
        self.assertEqual(annual_products["FUNDOS"]["meta_acumulada"], 120)
        self.assertEqual(annual_products["FUNDOS"]["producao_acumulada"], 20)
        self.assertEqual(annual_products["CREDITO_PESSOAL"]["meta_acumulada"], 110)
        self.assertEqual(annual_products["CREDITO_PESSOAL"]["producao_acumulada"], 105)

        period = analytics.accumulated_production(
            "2026-01", "2026-02", "SINGULAR", "NO_PERIODO"
        )
        period_products = {item["produto"]: item for item in period["products"]}
        self.assertEqual(period_products["FUNDOS"]["meta_acumulada"], 20)
        self.assertEqual(period_products["FUNDOS"]["producao_acumulada"], 20)
        self.assertEqual(period_products["CREDITO_PESSOAL"]["meta_acumulada"], 60)
        self.assertEqual(period_products["CREDITO_PESSOAL"]["producao_acumulada"], 55)

        single_month = analytics.accumulated_production(
            "2026-01", "2026-01", "SINGULAR", "NO_PERIODO"
        )
        single_products = {item["produto"]: item for item in single_month["products"]}
        self.assertEqual(single_products["FUNDOS"]["meta_acumulada"], 10)
        self.assertEqual(single_products["FUNDOS"]["producao_acumulada"], 8)

        analytic = analytics.monthly_product_evolution(
            "2026-01",
            "2026-02",
            "SINGULAR",
        )
        analytic_products = {
            item["produto"]: item
            for item in analytic["products"]
        }
        self.assertEqual(
            [point["meta"] for point in analytic_products["FUNDOS"]["points"]],
            [10, 10],
        )
        self.assertEqual(
            [point["producao"] for point in analytic_products["FUNDOS"]["points"]],
            [8, 12],
        )
        self.assertEqual(
            [point["meta"] for point in analytic_products["CREDITO_PESSOAL"]["points"]],
            [50, 60],
        )
        self.assertEqual(
            [point["producao"] for point in analytic_products["CREDITO_PESSOAL"]["points"]],
            [45, 55],
        )
        self.assertEqual(
            analytic_products["CREDITO_PESSOAL"]["regra_calculo"],
            "saldo_mensal_nao_acumulado",
        )

    def test_focus_product_performance_compares_agencies_in_selected_range(self) -> None:
        rows = []
        for period, pa_01_production, pa_02_production in (
            ("2026-01-01", 110, 70),
            ("2026-02-01", 100, 80),
        ):
            rows.extend(
                [
                    {
                        "MÊS/ANO": pd.Timestamp(period),
                        "PA": "01",
                        "PRODUTO": "COBRANCA",
                        "META_MENSAL": 100,
                        "PRODUCAO_MENSAL": pa_01_production,
                        "META_ANUAL": 200,
                        "PRODUCAO_ANUAL": pa_01_production,
                    },
                    {
                        "MÊS/ANO": pd.Timestamp(period),
                        "PA": "02",
                        "PRODUTO": "COBRANCA",
                        "META_MENSAL": 100,
                        "PRODUCAO_MENSAL": pa_02_production,
                        "META_ANUAL": 200,
                        "PRODUCAO_ANUAL": pa_02_production,
                    },
                ]
            )
        self.storage.frame = pd.DataFrame(rows)

        result = analytics.focus_product_agency_performance(
            "2026-01", "2026-02", "NO_PERIODO"
        )
        cobranca = next(
            product for product in result["products"]
            if product["produto"] == "COBRANCA"
        )
        units = {unit["pa"]: unit for unit in cobranca["units"]}

        self.assertEqual(result["range"]["start"], "2026-01")
        self.assertEqual(result["range"]["end"], "2026-02")
        self.assertEqual(cobranca["summary"]["units"], 2)
        self.assertEqual(cobranca["summary"]["adherent"], 1)
        self.assertEqual(units["01"]["meta_acumulada"], 200)
        self.assertEqual(units["01"]["producao_acumulada"], 210)
        self.assertEqual(units["01"]["status"], "aderente")
        self.assertEqual(
            [point["meta"] for point in units["01"]["points"]],
            [100, 100],
        )
        self.assertEqual(
            [point["producao"] for point in units["01"]["points"]],
            [110, 100],
        )
        self.assertAlmostEqual(
            units["01"]["points"][0]["percentual_atingido"],
            110,
        )
        self.assertEqual(units["01"]["points"][0]["gap"], 10)
        self.assertEqual(units["01"]["points"][0]["status"], "aderente")
        self.assertEqual(units["02"]["points"][0]["gap"], -30)
        self.assertEqual(units["02"]["points"][0]["status"], "nao_aderente")
        self.assertEqual(
            cobranca["monthly_summary"],
            {"expected": 4, "adherent": 2, "adherence_rate": 50.0},
        )
        self.assertEqual(cobranca["consolidated"]["period"]["meta_acumulada"], 400)
        self.assertEqual(cobranca["consolidated"]["period"]["producao_acumulada"], 360)
        self.assertEqual(cobranca["consolidated"]["annual"]["meta_acumulada"], 400)
        self.assertEqual(cobranca["consolidated"]["annual"]["producao_acumulada"], 180)
        self.assertEqual(units["02"]["producao_acumulada"], 150)
        self.assertEqual(units["02"]["status"], "nao_aderente")

    def test_product_evolution_study_calculates_required_year_end_pace(self) -> None:
        self.storage.frame = pd.DataFrame(
            [
                {
                    importer.PERIOD_COLUMN: pd.Timestamp(period),
                    "PA": "01",
                    "PRODUTO": "COBRANCA",
                    "META_MENSAL": 100,
                    "PRODUCAO_MENSAL": production,
                    "META_ANUAL": 1200,
                    "PRODUCAO_ANUAL": annual_production,
                }
                for period, production, annual_production in (
                    ("2026-01-01", 100, 100),
                    ("2026-02-01", 150, 250),
                    ("2026-03-01", 150, 400),
                )
            ]
        )

        result = analytics.product_evolution_study("01", "COBRANCA")

        self.assertEqual(result["reference_period"], "2026-03")
        self.assertEqual(result["unit"]["code"], "01")
        self.assertEqual(result["product"]["value"], "COBRANCA")
        self.assertEqual(result["summary"]["annual_target"], 1200)
        self.assertEqual(result["summary"]["current_production"], 400)
        self.assertEqual(result["summary"]["gap"], 800)
        self.assertEqual(result["summary"]["months_remaining"], 9)
        self.assertAlmostEqual(result["summary"]["required_per_month"], 800 / 9)
        self.assertAlmostEqual(result["summary"]["average_monthly_pace"], 400 / 3)
        self.assertAlmostEqual(result["summary"]["recommended_per_month"], 400 / 3)
        self.assertEqual(result["summary"]["current_month_production"], 150)
        self.assertEqual(result["summary"]["monthly_evolution"], 150)
        self.assertTrue(result["summary"]["will_reach_target"])
        self.assertEqual(result["summary"]["status"], "no_ritmo")
        self.assertEqual(len(result["points"]), 12)
        self.assertEqual(result["points"][1]["monthly_target"], 100)
        self.assertEqual(result["points"][1]["monthly_production"], 150)
        self.assertEqual(result["points"][1]["monthly_gap"], 50)
        self.assertEqual(result["points"][1]["accumulated_deficit"], 0)
        self.assertEqual(result["points"][1]["accumulated_surplus"], 50)
        self.assertIsNone(result["points"][2]["evolution_plan"])
        self.assertEqual(result["points"][3]["monthly_target"], 100)
        self.assertEqual(result["points"][3]["monthly_target_source"], "annual_reference")
        self.assertAlmostEqual(result["points"][3]["evolution_plan"], 400 / 3)
        self.assertAlmostEqual(result["points"][3]["plan_additional_to_base"], (400 / 3) - 100)
        self.assertEqual(result["points"][2]["recommended_projection"], 400)
        self.assertEqual(result["points"][-1]["recommended_projection"], 1600)

    def test_product_evolution_study_recommends_acceleration_when_projection_misses_target(self) -> None:
        self.storage.frame = pd.DataFrame(
            [
                {
                    importer.PERIOD_COLUMN: pd.Timestamp(period),
                    "PA": "01",
                    "PRODUTO": "COBRANCA",
                    "META_MENSAL": 100,
                    "PRODUCAO_MENSAL": 50,
                    "META_ANUAL": 1200,
                    "PRODUCAO_ANUAL": annual_production,
                }
                for period, annual_production in (
                    ("2026-01-01", 50),
                    ("2026-02-01", 100),
                    ("2026-03-01", 150),
                )
            ]
        )

        result = analytics.product_evolution_study("01", "COBRANCA")

        self.assertFalse(result["summary"]["will_reach_target"])
        self.assertEqual(result["summary"]["average_monthly_pace"], 50)
        self.assertAlmostEqual(result["summary"]["required_per_month"], 1050 / 9)
        self.assertAlmostEqual(result["summary"]["recommended_per_month"], 1050 / 9)
        self.assertAlmostEqual(result["summary"]["pace_increase_required"], (1050 / 9) - 50)
        self.assertEqual(result["summary"]["base_monthly_target"], 100)
        self.assertAlmostEqual(result["summary"]["plan_additional_to_base"], (1050 / 9) - 100)
        self.assertEqual(result["points"][2]["accumulated_deficit"], 150)
        self.assertAlmostEqual(result["points"][3]["evolution_plan"], 1050 / 9)
        self.assertAlmostEqual(result["points"][-1]["recommended_projection"], 1200)

    def test_evolution_deficit_accumulates_gaps_without_summing_balance_products(self) -> None:
        rows = []
        for period, regular_production, regular_annual, balance_annual in (
            ("2026-01-01", 80, 500, 50),
            ("2026-02-01", 90, 500, 120),
        ):
            rows.extend(
                [
                    {
                        importer.PERIOD_COLUMN: pd.Timestamp(period),
                        "PA": "01",
                        "PRODUTO": "COBRANCA",
                        "META_MENSAL": 100,
                        "PRODUCAO_MENSAL": regular_production,
                        "META_ANUAL": 1200,
                        "PRODUCAO_ANUAL": regular_annual,
                    },
                    {
                        importer.PERIOD_COLUMN: pd.Timestamp(period),
                        "PA": "01",
                        "PRODUTO": "CREDITO_PESSOAL",
                        "META_MENSAL": 100,
                        "PRODUCAO_MENSAL": 80,
                        "META_ANUAL": 1200,
                        "PRODUCAO_ANUAL": balance_annual,
                    },
                ]
            )
        self.storage.frame = pd.DataFrame(rows)

        regular = analytics.product_evolution_study("01", "COBRANCA")
        balance = analytics.product_evolution_study("01", "CREDITO_PESSOAL")

        self.assertEqual(
            [point["accumulated_deficit"] for point in regular["points"][:2]],
            [20, 30],
        )
        self.assertEqual(
            [point["accumulated_deficit"] for point in balance["points"][:2]],
            [50, 80],
        )

    def test_agency_cards_pdf_respects_period_and_contains_all_units(self) -> None:
        self.storage.frame = pd.DataFrame(
            [
                {
                    "MÊS/ANO": pd.Timestamp("2026-01-01"),
                    "PA": unit,
                    "PRODUTO": "COBRANCA",
                    "META_MENSAL": 100,
                    "PRODUCAO_MENSAL": production,
                    "META_ANUAL": 1200,
                    "PRODUCAO_ANUAL": production,
                }
                for unit, production in (("01", 120), ("04", 80), ("89", 500))
            ]
        )

        content, name = exporter.export_agency_cards_pdf("2026-01", "2026-01")

        self.assertIsInstance(content, bytes)
        self.assertTrue(name.startswith("acompanhamento_agencias_"))
        self.assertTrue(content.startswith(b"%PDF"))
        self.assertGreater(len(content), 1_000)


if __name__ == "__main__":
    unittest.main()
