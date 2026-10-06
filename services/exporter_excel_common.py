from __future__ import annotations

from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A3, landscape
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas

from app.config import EXPECTED_COLUMNS, NUMERIC_COLUMNS, PERIOD_COLUMN
from repo import database
from services.analytics import (
    accumulated_production,
    agency_accumulated_production,
    monthly_product_evolution,
    unit_performance_summary,
)
from services.importer import ImportValidationError


_HEADER_FILL = PatternFill("solid", fgColor="163E35")
_REPORT_FILL = PatternFill("solid", fgColor="003641")
_SECTION_FILL = PatternFill("solid", fgColor="003641")
_COLUMN_HEADER_FILL = PatternFill("solid", fgColor="00AE9D")
_SPACER_FILL = PatternFill("solid", fgColor="C9D200")
_ALT_ROW_FILL = PatternFill("solid", fgColor="F5F8F6")
_SUCCESS_FILL = PatternFill("solid", fgColor="E6F6EB")
_ATTENTION_FILL = PatternFill("solid", fgColor="FFF0ED")
_ZERO_FILL = PatternFill("solid", fgColor="F0F4F2")
_CELL_BORDER = Border(
    left=Side(style="thin", color="E1E9E5"),
    right=Side(style="thin", color="E1E9E5"),
    top=Side(style="thin", color="E1E9E5"),
    bottom=Side(style="thin", color="D5E0DB"),
)
_MONTH_NAMES = {
    1: "JANEIRO",
    2: "FEVEREIRO",
    3: "MARÇO",
    4: "ABRIL",
    5: "MAIO",
    6: "JUNHO",
    7: "JULHO",
    8: "AGOSTO",
    9: "SETEMBRO",
    10: "OUTUBRO",
    11: "NOVEMBRO",
    12: "DEZEMBRO",
}


def _style_sheet(sheet, frame: pd.DataFrame, freeze_panes: str = "A2") -> None:
    sheet.freeze_panes = freeze_panes
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, column in enumerate(frame.columns, start=1):
        values = frame[column].fillna("").astype(str)
        longest = int(values.str.len().max()) if not values.empty else 0
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(str(column)) + 2, longest + 2),
            34,
        )


def _is_unit_product(product: object) -> bool:
    return str(product).upper().startswith("PREVIDENCIA")


def _humanize_product(product: object) -> str:
    return str(product).replace("_", " ").title()


def _period_heading(period: str) -> str:
    value = pd.Period(period, freq="M")
    return f"{_MONTH_NAMES[value.month]}/{value.year}"


def _range_heading(data: dict) -> str:
    start = _period_heading(data["range"]["start"])
    end = _period_heading(data["range"]["end"])
    return start if start == end else f"{start} A {end}"


def _monthly_volume_frames(data: dict) -> list[tuple[str, pd.DataFrame]]:
    frames: list[tuple[str, pd.DataFrame]] = []
    for period in data["periods"]:
        rows: list[dict] = []
        for product in data["products"]:
            point = next(
                (item for item in product["points"] if item["period"] == period["value"]),
                None,
            )
            if point is None or point["meta"] is None:
                continue
            rows.append(
                {
                    "PRODUTO": _humanize_product(product["produto"]),
                    "META": point["meta"],
                    "PRODUÇÃO": point["producao"],
                    "%": point["percentual_atingido"],
                    "GAP": point["gap"],
                }
            )
        frames.append(
            (
                f"PRODUÇÃO EM VOLUME — {_period_heading(period['value'])}",
                pd.DataFrame(rows),
            )
        )
    return frames


def _consolidated_frame(data: dict, annual: bool = False) -> pd.DataFrame:
    meta_column = "META ANUAL" if annual else "META ACUMULADA"
    production_column = "PRODUÇÃO ACUMULADA"
    return pd.DataFrame(
        [
            {
                "PRODUTO": _humanize_product(product["produto"]),
                meta_column: product["meta_acumulada"],
                production_column: product["producao_acumulada"],
                "%": product["percentual_atingido"],
                "GAP": product["gap"],
            }
            for product in data["products"]
        ]
    )


def _product_monthly_score_frame(data: dict) -> pd.DataFrame:
    rows: list[dict] = []
    for product in data["products"]:
        row: dict[str, object] = {"PRODUTO": _humanize_product(product["produto"])}
        achieved = 0
        for point in product["points"]:
            monthly_score = int(
                point["percentual_atingido"] is not None
                and point["percentual_atingido"] >= 100
            )
            row[_period_heading(point["period"])] = monthly_score
            achieved += monthly_score
        target = len(product["points"])
        row["METAS NO PERÍODO"] = target
        row["METAS ATINGIDAS"] = achieved
        row["% APROVEITAMENTO"] = float(achieved / target * 100) if target else 0.0
        rows.append(row)
    return pd.DataFrame(rows)


def _unit_monthly_score_frame(data: dict) -> pd.DataFrame:
    rows: list[dict] = []
    for unit in sorted(data["units"], key=lambda item: str(item["code"]).zfill(3)):
        row: dict[str, object] = {"AGÊNCIA": unit["label"]}
        for month in unit["monthly"]:
            row[_period_heading(month["period"])] = month["adherent"]
        row["METAS NO PERÍODO"] = unit["products"]
        row["METAS ATINGIDAS"] = unit["adherent"]
        row["% APROVEITAMENTO"] = unit["adherence_rate"]
        rows.append(row)
    return pd.DataFrame(rows)


def _write_report_section(
    writer: pd.ExcelWriter,
    sheet_name: str,
    title: str,
    frame: pd.DataFrame,
    start_row: int,
    value_format: str,
) -> int:
    if frame.empty:
        raise ImportValidationError(f"Não existem dados para compor a seção: {title}.")

    frame.to_excel(writer, sheet_name=sheet_name, startrow=start_row + 1, index=False)
    sheet = writer.sheets[sheet_name]
    title_row = start_row + 1
    header_row = start_row + 2
    last_column = len(frame.columns)
    sheet.merge_cells(
        start_row=title_row,
        start_column=1,
        end_row=title_row,
        end_column=last_column,
    )
    title_cell = sheet.cell(title_row, 1, title)
    title_cell.fill = _SECTION_FILL
    title_cell.font = Font(name="Aptos", color="FFFFFF", bold=True, size=11)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[title_row].height = 27

    for cell in sheet[header_row]:
        if cell.column > last_column:
            break
        cell.fill = _COLUMN_HEADER_FILL
        cell.font = Font(name="Aptos", color="032F2A", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _CELL_BORDER
    sheet.row_dimensions[header_row].height = 25

    first_data_row = header_row + 1
    last_data_row = header_row + len(frame)
    for row_index in range(first_data_row, last_data_row + 1):
        is_unit_product = _is_unit_product(frame.iloc[row_index - first_data_row][frame.columns[0]])
        for column_index, column in enumerate(frame.columns, start=1):
            cell = sheet.cell(row_index, column_index)
            cell.border = _CELL_BORDER
            cell.font = Font(name="Aptos", color="233D37", size=10)
            if (row_index - first_data_row) % 2:
                cell.fill = _ALT_ROW_FILL
            if column_index == 1:
                cell.alignment = Alignment(horizontal="left")
                cell.font = Font(name="Aptos", color="173A32", bold=True, size=10)
            else:
                cell.alignment = Alignment(horizontal="center")
            if column in {"%", "% APROVEITAMENTO"}:
                cell.number_format = '0.0"%"'
                if column == "% APROVEITAMENTO":
                    cell.font = Font(name="Aptos", color="173A32", bold=True, size=10)
                elif isinstance(cell.value, (int, float)) and cell.value >= 100:
                    cell.fill = _SUCCESS_FILL
                    cell.font = Font(name="Aptos", color="167357", bold=True, size=10)
                elif isinstance(cell.value, (int, float)):
                    cell.fill = _ATTENTION_FILL
                    cell.font = Font(name="Aptos", color="C43F2D", bold=True, size=10)
            elif value_format == "volume" and isinstance(cell.value, (int, float)):
                cell.number_format = (
                    '#,##0;[Red]-#,##0'
                    if is_unit_product
                    else 'R$ #,##0;[Red]-R$ #,##0'
                )
                if column == "GAP":
                    color = "167357" if cell.value >= 0 else "C43F2D"
                    cell.font = Font(name="Aptos", color=color, bold=True, size=10)
            elif value_format == "score" and isinstance(cell.value, (int, float)):
                cell.number_format = "0"
                if column in {"METAS NO PERÍODO", "METAS ATINGIDAS"}:
                    cell.font = Font(name="Aptos", color="173A32", bold=True, size=10)
                elif cell.value > 0:
                    cell.fill = _SUCCESS_FILL
                    cell.font = Font(name="Aptos", color="167357", bold=True, size=10)
                else:
                    cell.fill = _ZERO_FILL
                    cell.font = Font(name="Aptos", color="81918B", size=10)
        sheet.row_dimensions[row_index].height = 22

    for column_index, column in enumerate(frame.columns, start=1):
        values = frame[column].fillna("").astype(str)
        longest = int(values.str.len().max()) if not values.empty else 0
        minimum = 25 if column_index == 1 else 14
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(minimum, len(str(column)) + 2, longest + 2),
            36,
        )
    if "% APROVEITAMENTO" in frame.columns:
        percentage_column = frame.columns.get_loc("% APROVEITAMENTO") + 1
        percentage_letter = get_column_letter(percentage_column)
        sheet.conditional_formatting.add(
            f"{percentage_letter}{first_data_row}:{percentage_letter}{last_data_row}",
            ColorScaleRule(
                start_type="min",
                start_color="F8696B",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFEB84",
                end_type="max",
                end_color="63BE7B",
            ),
        )
    return last_data_row + 2


def _style_report_sheet(
    sheet,
    report_title: str,
    period_label: str,
    generated_label: str,
    tab_color: str,
) -> None:
    last_column = max(sheet.max_column, 5)
    for row_index in range(1, 4):
        for column_index in range(1, last_column + 1):
            sheet.cell(row_index, column_index).fill = _REPORT_FILL
    for column_index in range(1, last_column + 1):
        sheet.cell(4, column_index).fill = _SPACER_FILL

    for row_index in range(1, 4):
        sheet.merge_cells(
            start_row=row_index,
            start_column=1,
            end_row=row_index,
            end_column=last_column,
        )
    sheet["A1"] = "ACOMPANHAMENTO DE PRODUTOS"
    sheet["A1"].font = Font(name="Aptos Display", color="C9D200", bold=True, size=10)
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet["A2"] = report_title
    sheet["A2"].font = Font(name="Aptos Display", color="FFFFFF", bold=True, size=19)
    sheet["A2"].alignment = Alignment(horizontal="left", vertical="center")
    sheet["A3"] = f"PERÍODO: {period_label}   |   EMITIDO EM: {generated_label}"
    sheet["A3"].font = Font(name="Aptos", color="B9D5CC", size=9)
    sheet["A3"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 20
    sheet.row_dimensions[2].height = 31
    sheet.row_dimensions[3].height = 21
    sheet.row_dimensions[4].height = 4

    sheet.sheet_properties.tabColor = tab_color
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.freeze_panes = None
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.55
    sheet.page_margins.bottom = 0.55
    sheet.page_margins.header = 0.2
    sheet.page_margins.footer = 0.2
    sheet.oddHeader.center.text = "&B ACOMPANHAMENTO DE PRODUTOS"
    sheet.oddHeader.center.size = 8
    sheet.oddHeader.center.color = "55736A"
    sheet.oddFooter.left.text = sheet.title
    sheet.oddFooter.center.text = "Página &P de &N"
    sheet.oddFooter.right.text = generated_label
    sheet.oddFooter.left.size = 8
    sheet.oddFooter.center.size = 8
    sheet.oddFooter.right.size = 8
    sheet.print_area = f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    sheet.auto_filter.ref = None
